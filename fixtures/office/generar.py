"""Genera los archivos de prueba de Office (datos ficticios). Uso: python generar.py"""
import io
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def logo_png(color, size=(120, 60)):
    img = Image.new("RGB", size, color)
    ImageDraw.Draw(img).rectangle([8, 8, size[0] - 8, size[1] - 8], outline="white", width=3)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos generales"
    thin = Side(style="thin", color="444444")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    answer = PatternFill("solid", fgColor="FFF2CC")
    head = PatternFill("solid", fgColor="1F4E78")

    ws.add_image(XlImage(logo_png("#1F4E78")), "A1")
    ws.merge_cells("C1:F2")
    ws["C1"] = "CUESTIONARIO DE SOLICITUD (FICTICIO)"
    ws["C1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["C1"].fill = head
    ws["C1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 30
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 18

    rows = [
        (5, "Nombre de la institución"),
        (6, "Número de beneficiarios directos"),
        (7, "Fecha de inicio del proyecto"),
        (8, "Tipo de institución"),
        (9, "Monto solicitado (MXN)"),
        (10, "Monto aportado por la institución (MXN)"),
    ]
    for r, q in rows:
        ws.cell(r, 1, q).font = Font(bold=True)
        ws.cell(r, 1).border = box
        ws.cell(r, 2).fill = answer
        ws.cell(r, 2).border = box
    ws.merge_cells("B5:D5")  # respuesta en celda combinada
    ws["B9"].number_format = '"$"#,##0.00'
    ws["B10"].number_format = '"$"#,##0.00'
    ws["B7"].number_format = "dd/mm/yyyy"

    ws["A11"] = "Total del proyecto (MXN)"
    ws["A11"].font = Font(bold=True)
    ws["B11"] = "=B9+B10"
    ws["B11"].number_format = '"$"#,##0.00'
    ws["A12"] = "Porcentaje de aportación propia"
    ws["B12"] = "=IF(B11=0,0,B10/B11)"
    ws["B12"].number_format = "0.0%"

    ws["A14"] = "Describa el impacto esperado"
    ws["A14"].font = Font(bold=True)
    ws.merge_cells("A15:F18")
    ws["A15"].fill = answer
    ws["A15"].alignment = Alignment(wrap_text=True, vertical="top")

    dv = DataValidation(type="list", formula1='"Asilo,Casa hogar,Otra"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B8")

    ws2 = wb.create_sheet("Presupuesto")
    ws2.add_image(XlImage(logo_png("#C55A11", (80, 40))), "F1")
    ws2.append(["Concepto", "Cantidad", "Precio unitario", "Importe"])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
    for i in range(2, 8):
        ws2.cell(i, 1).fill = answer
        ws2.cell(i, 2).fill = answer
        ws2.cell(i, 3).fill = answer
        ws2.cell(i, 4, f"=B{i}*C{i}")
    ws2["A8"] = "TOTAL"
    ws2["D8"] = "=SUM(D2:D7)"
    ws2["A8"].font = Font(bold=True)

    ws3 = wb.create_sheet("Catálogos")
    ws3["A1"] = "uso interno"
    ws3.sheet_state = "hidden"
    wb.save("cuestionario-prueba.xlsx")


def word():
    d = Document()
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Convocatoria ficticia · Encabezado"
    sec.header.paragraphs[0].runs[0].font.bold = True
    sec.header.paragraphs[0].add_run().add_picture(logo_png("#1F4E78", (120, 40)), width=Inches(1.2))
    sec.footer.paragraphs[0].text = "Pie de página · {{institucion_nombre}}"
    sec.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    d.add_heading("Proyecto: {{titulo_proyecto}}", level=1)

    d.add_heading("1. Justificación", level=2)
    d.add_paragraph("{{section:justificacion}}")

    d.add_heading("2. Población beneficiada", level=2)
    p = d.add_paragraph()
    r = p.add_run("Atendemos a ")
    r2 = p.add_run("{{beneficiarios")  # marcador partido en dos runs (caso difícil)
    r2.bold = True
    r2.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    p.add_run("_total}}").bold = True
    p.add_run(" personas.")

    d.add_heading("3. Objetivo", level=2)
    d.add_paragraph("{{section:objetivo}}", style="List Bullet")

    d.add_heading("4. Presupuesto", level=2)
    d.add_paragraph("{{table:presupuesto}}")
    t = d.add_table(rows=2, cols=4)
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(["Concepto", "Cantidad", "Precio", "Importe"]):
        t.cell(0, i).text = h
    t.cell(1, 0).text = "(ejemplo existente)"

    d.add_heading("5. Responsable", level=2)
    d.add_paragraph("Elaboró: {{institucion_nombre}}")
    d.core_properties.author = "Autor de prueba (debe limpiarse)"
    d.save("plantilla-prueba.docx")


def pdf():
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak

    st = getSampleStyleSheet()
    doc = SimpleDocTemplate("convocatoria-con-tablas.pdf", pagesize=letter)
    t = Table(
        [["Concepto", "Monto máximo (MXN)", "Plazo (meses)"],
         ["Obra y equipamiento", "$300,000", "12"],
         ["Capacitación", "$50,000", "6"]]
    )
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
    doc.build([
        Paragraph("Convocatoria ficticia 2026", st["Title"]),
        Paragraph("Requisito REQ-01: el monto solicitado no puede exceder los límites de la tabla.", st["Normal"]),
        Spacer(1, 12), t, PageBreak(),
        Paragraph("Documentos requeridos: acta constitutiva, comprobante de domicilio, cotizaciones.", st["Normal"]),
    ])
    # "scanned" PDF: pages are only images, no text layer
    img = Image.new("RGB", (850, 1100), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 80), "Convocatoria escaneada ficticia", fill="black")
    d.rectangle([60, 140, 700, 400], outline="black", width=2)
    img.save("convocatoria-escaneada.pdf", "PDF", resolution=100.0)


def scans():
    """Scanned pages with a known fictitious text, to read them with OCR (Tesseract)."""
    import random
    from PIL import ImageFilter, ImageFont

    font = ImageFont.truetype("C:/Windows/Fonts/times.ttf", 34)
    bold = ImageFont.truetype("C:/Windows/Fonts/timesbd.ttf", 40)

    def page(title, lines):
        # an A4 sheet at 200 dpi, a little off white, as a scanner leaves it
        img = Image.new("L", (1654, 2339), 246)
        d = ImageDraw.Draw(img)
        d.text((160, 180), title, fill=20, font=bold)
        y = 300
        for line in lines:
            d.text((160, y), line, fill=25, font=font)
            y += 52
        return img

    def worn(img):
        # a bad copy: blurred, washed out and dirty
        img = img.filter(ImageFilter.GaussianBlur(1.7)).point(lambda v: 130 + v * 120 // 255)
        rnd = random.Random(7)
        px = img.load()
        for _ in range(20000):
            x, y = rnd.randrange(img.width), rnd.randrange(img.height)
            px[x, y] = rnd.randrange(90, 250)
        return img

    first = page("ACTA CONSTITUTIVA", [
        "Casa Hogar Ficticia del Valle, I.A.P.",
        "",
        "En la Ciudad de México, a los doce días del mes de marzo de dos mil",
        "quince, se reunieron los asociados fundadores para constituir una",
        "institución de asistencia privada sin fines de lucro.",
        "",
        "Objeto social: brindar albergue, alimentación y cuidado a personas",
        "adultas mayores en situación de abandono.",
        "",
        "Domicilio social: Calle de los Robles número 45, colonia Centro.",
        "Para dudas sobre este documento: contacto@casaficticia.example.org",
    ])
    second = worn(page("CLÁUSULA SEGUNDA", [
        "El patrimonio de la institución se integra con las aportaciones",
        "de los asociados y con los donativos que reciba para su objeto.",
        "Ningún asociado podrá disponer de ese patrimonio para sí.",
    ]))
    first.save("acta-escaneada.pdf", "PDF", resolution=200.0, save_all=True, append_images=[second])

    # a photo of a page taken with a phone: a little turned, on a gray table
    sheet = page("CONSTANCIA DE REGISTRO", [
        "La Casa Hogar Ficticia del Valle quedó inscrita en el padrón",
        "de instituciones de asistencia social de este año.",
        "Vigente hasta el treinta y uno de diciembre.",
    ]).crop((0, 0, 1654, 900)).rotate(1.5, expand=True, fillcolor=120)
    sheet.convert("RGB").save("constancia-foto.jpg", "JPEG", quality=85)

    # an invoice: a table of concepts, and the email of who sells
    sans = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 30)
    inv = Image.new("L", (1654, 800), 250)
    d = ImageDraw.Draw(inv)
    d.text((120, 70), "FACTURA", fill=20, font=bold)
    d.text((120, 150), "Proveedor: Abarrotes Ficticios del Centro", fill=20, font=sans)
    d.text((120, 200), "Correo: ventas@abarrotesficticios.mx", fill=20, font=sans)
    rows = [
        ["Concepto", "Cantidad", "Precio", "Importe"],
        ["Arroz", "10", "$28.00", "$280.00"],
        ["Frijol", "5", "$35.00", "$175.00"],
        ["Aceite", "4", "$45.00", "$180.00"],
    ]
    for i, row in enumerate(rows):
        for x, cell in zip([120, 700, 1000, 1300], row):
            d.text((x, 300 + i * 60), cell, fill=20, font=sans)
    d.text((1000, 580), "Total", fill=20, font=sans)
    d.text((1300, 580), "$635.00", fill=20, font=sans)
    inv.convert("RGB").save("factura-foto.jpg", "JPEG", quality=88)


def small():
    """Pictures with small letters: a screenshot and a little photo, read well only once enlarged (ADR-035)."""
    from PIL import ImageFilter, ImageFont

    sans = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 10)
    sans_bold = ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", 11)

    # a screenshot of the record of a public registry, at 96 dpi: letters about 7 pixels tall
    shot = Image.new("RGB", (520, 230), "white")
    d = ImageDraw.Draw(shot)
    d.rectangle([0, 0, 520, 24], fill=(31, 78, 120))
    d.text((10, 6), "Registro de Instituciones de Asistencia (ficticio)", fill="white", font=sans_bold)
    lines = [
        ("Folio de registro:", "48213"),
        ("Denominación:", "Casa Hogar Ficticia del Valle"),
        ("Representante:", "Lucía Hernández Ramírez"),
        ("Modalidad:", "Internado"),
        ("Población:", "personas adultas mayores"),
        ("Servicios:", "atención médica, alimentación y albergue"),
        ("Correo:", "registro@casaficticia.example.org"),
    ]
    for i, (label, value) in enumerate(lines):
        y = 38 + i * 26
        d.text((14, y), label, fill=(60, 60, 60), font=sans_bold)
        d.text((150, y), value, fill=(20, 20, 20), font=sans)
    shot.save("captura-chica.png", "PNG")

    # a little photo of a notice taken from afar: small, a little turned and soft, on a gray table. Read as it
    # comes, Tesseract finds no word in it at all
    note_font = ImageFont.truetype("C:/Windows/Fonts/times.ttf", 11)
    note = Image.new("L", (363, 126), 238)
    d = ImageDraw.Draw(note)
    for i, line in enumerate([
        "AVISO DE FUNCIONAMIENTO",
        "Casa Hogar Ficticia del Valle",
        "Horario de visitas: martes a domingo",
        "Responsable sanitario: Ernesto Martínez",
    ]):
        d.text((13, (14 + i * 30) * 11 // 13), line, fill=30, font=note_font)
    photo = note.rotate(1.2, expand=True, fillcolor=130).filter(ImageFilter.GaussianBlur(0.7))
    photo.convert("RGB").save("foto-chica.jpg", "JPEG", quality=70)


def protected():
    """A PDF that asks for a password to be opened: the program says so instead of calling it damaged."""
    from reportlab.pdfgen import canvas

    c = canvas.Canvas("protegido.pdf", encrypt="clave-ficticia")
    c.drawString(72, 720, "Documento ficticio con contrasena.")
    c.save()


def ruled():
    """A table drawn with lines inside a nested form, over a page background: what Pdfium reads as a grid (ADR-037)."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas("tabla-con-lineas.pdf", pagesize=letter)
    # a background that covers the page is design, not a table
    c.setFillGray(0.97)
    c.rect(0, 0, letter[0], letter[1], fill=1, stroke=0)
    c.setFillGray(0)
    c.drawString(72, 740, "Montos de la convocatoria ficticia")
    # the table lives in a form, drawn where the page puts it
    c.beginForm("tabla")
    rows = [("Concepto", "Monto"), ("Obra y equipamiento", "$300,000"), ("Capacitación", "$50,000")]
    for i, (a, b) in enumerate(rows):
        y = 100 - i * 24
        c.drawString(8, y + 7, a)
        c.drawString(208, y + 7, b)
    # Word draws borders as thin filled rectangles; others stroke them
    for i in range(4):
        c.rect(0, 124 - i * 24, 300, 0.5, fill=1, stroke=0)
    for x in (0, 200, 300):
        c.line(x, 52, x, 124)
    c.endForm()
    c.saveState()
    c.translate(72, 560)
    c.doForm("tabla")
    c.restoreState()
    c.showPage()
    c.save()


def mixed():
    """A page with a printed header and its body scanned: read as a picture, all of it (ADR-037)."""
    import io as _io
    from PIL import ImageFont
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    font = ImageFont.truetype("C:/Windows/Fonts/times.ttf", 34)
    body = Image.new("L", (1500, 900), 248)
    d = ImageDraw.Draw(body)
    for i, line in enumerate([
        "En la Ciudad de México se reunieron los asociados",
        "de la Casa Hogar Ficticia del Valle para aprobar",
        "el informe anual y el presupuesto del siguiente año.",
    ]):
        d.text((60, 80 + i * 60), line, fill=25, font=font)
    buf = _io.BytesIO()
    body.save(buf, "PNG")
    buf.seek(0)
    c = canvas.Canvas("pagina-mixta.pdf", pagesize=letter)
    c.drawString(72, 750, "Notaría Ficticia 12")
    c.drawImage(ImageReader(buf), 36, 150, width=540, height=560)
    c.showPage()
    c.save()


import sys

if sys.argv[1:] == ["escaneos"]:
    # only the scanned pages: the other files stay as they are
    scans()
elif sys.argv[1:] == ["chicas"]:
    # only the pictures with small letters
    small()
elif sys.argv[1:] == ["protegidos"]:
    # only the files with a password
    protected()
elif sys.argv[1:] == ["tablas"]:
    # only the table drawn with lines
    ruled()
elif sys.argv[1:] == ["mixta"]:
    # only the page with a printed header and a scanned body
    mixed()
else:
    excel()
    word()
    pdf()
    scans()
    small()
    protected()
    ruled()
    mixed()
print("ok")
